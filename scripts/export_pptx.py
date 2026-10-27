"""Export the LaTeX decks for submission: the PDF as is, plus a PowerPoint with one full-slide
300 dpi image per PDF page (not editable).

    python scripts/export_pptx.py     # deck/sentinel.pdf -> presentation/Sentinel.pdf + Sentinel.pptx
                                      # deck/backup.pdf   -> presentation/Sentinel_backup.pdf + .pptx
Build the PDFs first (see deck/sentinel.tex). Close the files in PowerPoint before running.
"""
import io
import shutil
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
DECK, OUT = ROOT / "deck", ROOT / "presentation"
PAIRS = (("sentinel.pdf", "Sentinel"), ("backup.pdf", "Sentinel_backup"))


def export(pdf, stem):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9
    for page in fitz.open(pdf):
        png = io.BytesIO(page.get_pixmap(dpi=300).tobytes("png"))
        slide = prs.slides.add_slide(prs.slide_layouts[6])               # blank
        slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    shutil.copyfile(pdf, OUT / f"{stem}.pdf")
    target = OUT / f"{stem}.pptx"
    try:
        prs.save(target)
    except PermissionError:                  # open in PowerPoint: Windows locks the file
        target = OUT / f"{stem}_new.pptx"
        prs.save(target)
        print(f"  {stem}.pptx is open in PowerPoint; saved {target.name} instead (close it and re-run to replace)")
    print(f"wrote presentation/{stem}.pdf + {target.name}: {len(prs.slides)} slides")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for pdf, stem in PAIRS:
        export(DECK / pdf, stem)
