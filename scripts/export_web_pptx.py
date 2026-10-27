"""Export the web deck for submission: one full-slide 3840x2160 image per slide.

    python scripts/export_web_pptx.py            # web/ -> presentation/Sentinel_web.pptx + .pdf (original colours)
    python scripts/export_web_pptx.py dark       # -> presentation/Sentinel_web_dark.pptx + .pdf

The PPTX carries the speaker notes. Slides are images (not editable): edit the text in
web/index.html and re-run. Close the files in PowerPoint / the PDF viewer before running.
"""
import io
import sys
from pathlib import Path

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
THEME = sys.argv[1] if len(sys.argv) > 1 else "light"
OUT = ROOT / "presentation"
STEM = "Sentinel_web" if THEME == "light" else f"Sentinel_web_{THEME}"


def main():
    url = (ROOT / "web" / "index.html").as_uri() + f"?still&theme={THEME}#1"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)   # 16:9
    pages = []
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome", headless=True)
        page = browser.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=2)
        page.goto(url)
        page.evaluate("document.fonts.ready")
        page.wait_for_timeout(800)
        n = page.evaluate("document.querySelectorAll('.slide').length")
        for i in range(1, n + 1):
            page.evaluate(f"location.hash = '#{i}'")
            page.wait_for_timeout(400)
            png = page.screenshot(type="png")
            notes = page.evaluate(f"document.querySelectorAll('.slide')[{i - 1}].querySelector('.notes')?.textContent.trim() || ''")
            slide = prs.slides.add_slide(prs.slide_layouts[6])               # blank
            slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
            slide.notes_slide.notes_text_frame.text = notes
            pages.append(Image.open(io.BytesIO(png)).convert("RGB"))
        browser.close()
    OUT.mkdir(exist_ok=True)
    prs.save(OUT / f"{STEM}.pptx")
    # 3840 px across 13.333 in = 288 dpi, so each PDF page is a 16:9 slide
    pages[0].save(OUT / f"{STEM}.pdf", save_all=True, append_images=pages[1:], resolution=288, quality=92)
    for ext in ("pptx", "pdf"):
        f = OUT / f"{STEM}.{ext}"
        print(f"wrote {f.relative_to(ROOT)}: {len(pages)} slides, {f.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
